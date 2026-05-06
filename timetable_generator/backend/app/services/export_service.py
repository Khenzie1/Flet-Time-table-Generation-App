"""Export service for timetables."""
import csv
import io
from typing import List, Dict
import uuid
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from fastapi.responses import StreamingResponse

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.models.timetable import Timetable, TimetableSlot
from app.models.teacher import Teacher
from app.models.subject import Subject
from app.models.class_ import Class
from app.models.room import Room


class ExportService:
    """Service for exporting timetables."""
    
    def __init__(self, db: AsyncSession):
        self.db = db
    
    async def export_to_csv(self, timetable_id: uuid.UUID) -> StreamingResponse:
        """Export timetable to CSV format."""
        # Load timetable data
        slots_by_class = await self._load_timetable_data(timetable_id)
        
        # Create CSV in memory
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write header
        days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
        writer.writerow(["Period"] + days)
        
        # Get all period numbers
        all_periods = set()
        for class_slots in slots_by_class.values():
            for slot in class_slots:
                all_periods.add(slot.period_number)
        
        # Write data for each period
        for period in sorted(all_periods):
            row = [f"Period {period}"]
            
            for day in days:
                cell_content = ""
                for class_name, slots in slots_by_class.items():
                    for slot in slots:
                        if slot.day_of_week == day and slot.period_number == period:
                            cell_content += f"{slot.subject_name} - {slot.teacher_name}\n"
                
                row.append(cell_content.strip())
            
            writer.writerow(row)
        
        output.seek(0)
        
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment;filename=timetable_{timetable_id}.csv"}
        )
    
    async def export_to_excel(self, timetable_id: uuid.UUID) -> StreamingResponse:
        """Export timetable to Excel format."""
        # Load timetable data
        slots_by_class = await self._load_timetable_data(timetable_id)
        
        # Create workbook
        wb = openpyxl.Workbook()
        
        # Define styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal='center', vertical='center')
        
        # Create sheet for each class
        days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
        
        for class_name, slots in slots_by_class.items():
            # Create sheet
            sheet_name = class_name[:31]  # Excel sheet name limit
            ws = wb.create_sheet(title=sheet_name)
            
            # Set headers
            ws['A1'] = "Period"
            ws['A1'].font = header_font
            ws['A1'].fill = header_fill
            ws['A1'].border = border
            ws['A1'].alignment = center_alignment
            
            for col, day in enumerate(days, start=2):
                cell = ws.cell(row=1, column=col, value=day)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = center_alignment
            
            # Get all period numbers for this class
            periods = sorted(set(s.period_number for s in slots))
            
            # Fill data
            for row, period in enumerate(periods, start=2):
                period_cell = ws.cell(row=row, column=1, value=f"Period {period}")
                period_cell.border = border
                period_cell.alignment = center_alignment
                
                for col, day in enumerate(days, start=2):
                    cell_content = ""
                    for slot in slots:
                        if slot.day_of_week == day and slot.period_number == period:
                            cell_content = f"{slot.subject_name}\n{slot.teacher_name}"
                            if slot.room_name:
                                cell_content += f"\n{slot.room_name}"
                    
                    cell = ws.cell(row=row, column=col, value=cell_content)
                    cell.border = border
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
            # Adjust column widths
            for col in range(1, len(days) + 2):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
            
            # Adjust row heights
            for row in range(1, len(periods) + 2):
                ws.row_dimensions[row].height = 60
        
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment;filename=timetable_{timetable_id}.xlsx"}
        )
    
    async def _load_timetable_data(self, timetable_id: uuid.UUID) -> Dict[str, List]:
        """Load timetable data organized by class."""
        # Load timetable with slots
        result = await self.db.execute(
            select(Timetable)
            .where(Timetable.id == timetable_id)
            .options(selectinload(Timetable.slots))
        )
        timetable = result.scalar_one_or_none()
        
        if not timetable:
            return {}
        
        # Load related data
        class_ids = list(set(s.class_id for s in timetable.slots))
        teacher_ids = list(set(s.teacher_id for s in timetable.slots))
        subject_ids = list(set(s.subject_id for s in timetable.slots))
        room_ids = list(set(s.room_id for s in timetable.slots if s.room_id))
        
        # Fetch related entities
        classes = {}
        if class_ids:
            result = await self.db.execute(
                select(Class).where(Class.id.in_(class_ids))
            )
            for cls in result.scalars().all():
                classes[cls.id] = cls
        
        teachers = {}
        if teacher_ids:
            result = await self.db.execute(
                select(Teacher).where(Teacher.id.in_(teacher_ids))
            )
            for teacher in result.scalars().all():
                teachers[teacher.id] = teacher
        
        subjects = {}
        if subject_ids:
            result = await self.db.execute(
                select(Subject).where(Subject.id.in_(subject_ids))
            )
            for subject in result.scalars().all():
                subjects[subject.id] = subject
        
        rooms = {}
        if room_ids:
            result = await self.db.execute(
                select(Room).where(Room.id.in_(room_ids))
            )
            for room in result.scalars().all():
                rooms[room.id] = room
        
        # Organize slots by class
        slots_by_class = {}
        for slot in timetable.slots:
            class_name = classes.get(slot.class_id, Class(class_name="Unknown")).class_name
            
            if class_name not in slots_by_class:
                slots_by_class[class_name] = []
            
            # Add enriched slot data
            slots_by_class[class_name].append(
                SlotExportData(
                    day_of_week=slot.day_of_week,
                    period_number=slot.period_number,
                    subject_name=subjects.get(slot.subject_id, Subject(name="Unknown")).name,
                    teacher_name=teachers.get(slot.teacher_id, Teacher(full_name="Unknown")).full_name,
                    room_name=rooms.get(slot.room_id, Room(room_name="")).room_name if slot.room_id else ""
                )
            )
        
        return slots_by_class


class SlotExportData:
    """Simple data class for export slots."""
    def __init__(self, day_of_week, period_number, subject_name, teacher_name, room_name):
        self.day_of_week = day_of_week
        self.period_number = period_number
        self.subject_name = subject_name
        self.teacher_name = teacher_name
        self.room_name = room_name