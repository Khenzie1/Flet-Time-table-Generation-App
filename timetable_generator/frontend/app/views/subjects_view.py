"""Subjects management view - Complete working version."""
import flet as ft
import uuid
import traceback

def _norm_role(role: str) -> str:
    """Normalise role string to uppercase regardless of source."""
    return (role or "VIEWER").strip().upper()

def _can_edit(role: str) -> bool:
    """Admin and Coordinator may mutate data; Viewer is read-only."""
    return _norm_role(role) in ("ADMIN", "COORDINATOR")



class SubjectsView:
    """Subjects management screen."""
    
    def __init__(self, app):
        self.app = app
        self.page = app.page
        self._role = _norm_role((app.current_user or {}).get("role", "VIEWER"))
        self.search_field = ft.TextField(
            hint_text="Search subjects...",
            prefix_icon=ft.icons.SEARCH,
            expand=True,
            on_change=self.search_subjects
        )
        self.category_tabs = ft.Tabs(
            selected_index=0,
            tabs=[
                ft.Tab(text="All"),
                ft.Tab(text="Core"),
                ft.Tab(text="Elective"),
                ft.Tab(text="Practical"),
            ],
            on_change=self.filter_by_category
        )
        self.subjects_table = ft.DataTable(
            columns=[
                ft.DataColumn(ft.Text("",             size=13)),
                ft.DataColumn(ft.Text("Code",         size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Subjects",         size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Category",     size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Periods/Week", size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Duration",     size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Requires Lab", size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
                ft.DataColumn(ft.Text("Actions",      size=13, weight=ft.FontWeight.W_600, color="#0F1E3A")),
            ],
            heading_row_color="#EFF6FF",
            heading_row_height=52,
            data_row_min_height=52,
            data_row_max_height=52,
            column_spacing=20,
            horizontal_margin=0,
            divider_thickness=1,
            show_checkbox_column=False,
            rows=[]
        )

        self.selected_ids = set()

        self._sel_count_lbl = ft.Text(
            "0 selected", size=14, weight=ft.FontWeight.BOLD, color=ft.colors.WHITE)
        self._sel_bar = ft.Container(
            content=ft.Row([
                ft.Icon(ft.icons.CHECK_BOX_OUTLINED, color=ft.colors.WHITE, size=20),
                self._sel_count_lbl,
                ft.Container(expand=True),
                ft.TextButton("Clear selection",
                    style=ft.ButtonStyle(color=ft.colors.WHITE70),
                    on_click=self._clear_selection),
                ft.ElevatedButton(
                    "Delete Selected",
                    icon=ft.icons.DELETE_SWEEP,
                    on_click=self._bulk_delete_confirm,
                    style=ft.ButtonStyle(
                        bgcolor=ft.colors.RED_700,
                        color=ft.colors.WHITE,
                        shape=ft.RoundedRectangleBorder(radius=8),
                    ),
                ),
            ], spacing=12),
            bgcolor="#0F1E3A",
            border_radius=10,
            padding=ft.padding.symmetric(horizontal=20, vertical=10),
            visible=False,
            margin=ft.margin.only(bottom=10),
        )
        
        # Add/Edit Dialog — save button stored for animation
        self._btn_save_subject = ft.ElevatedButton(
            "Save",
            icon=ft.icons.SAVE_ROUNDED,
            on_click=self._save_subject_animated,
            style=ft.ButtonStyle(
                bgcolor={"": "#2563EB", "hovered": "#1E40AF"},
                color=ft.colors.WHITE,
                shape=ft.RoundedRectangleBorder(radius=8),
                elevation={"": 0, "hovered": 2},
            ),
        )

        # Add/Edit Dialog
        self.dialog = ft.AlertDialog(
            title=ft.Text("Add Subject"),
            content=ft.Container(
                width=400,
                height=500,
                content=ft.Column([
                    ft.TextField(label="Subject Code", autofocus=True),
                    ft.TextField(label="Subject Name"),
                    ft.Dropdown(
                        label="Category",
                        options=[
                            ft.dropdown.Option("Core"),
                            ft.dropdown.Option("Elective"),
                            ft.dropdown.Option("Practical"),
                        ],
                        value="Core"
                    ),
                    ft.TextField(
                        label="Periods per Week",
                        value="3",
                        keyboard_type=ft.KeyboardType.NUMBER
                    ),
                    ft.TextField(
                        label="Duration (minutes)",
                        value="40",
                        keyboard_type=ft.KeyboardType.NUMBER
                    ),
                    ft.Checkbox(label="Requires Laboratory", value=False),
                ], spacing=10, scroll=ft.ScrollMode.AUTO),
                padding=20
            ),
            actions=[
                ft.TextButton("Cancel", on_click=self.close_dialog,
                              style=ft.ButtonStyle(color=ft.colors.GREY_600)),
                self._btn_save_subject,
            ]
        )
        
        self.current_subject_id = None
        self.all_subjects = []

        # Refresh button — stored so refresh() can animate it
        self._btn_refresh = ft.IconButton(
            icon=ft.icons.REFRESH_ROUNDED,
            icon_color=ft.colors.WHITE, icon_size=18,
            tooltip="Refresh",
            on_click=self.refresh,
        )

        # Live count labels — updated by display_subjects() on every data change
        self._lbl_count_top   = ft.Text(
            "0 subjects", size=12,
            color=ft.colors.with_opacity(0.6, ft.colors.WHITE))
        self._lbl_count_table = ft.Text(
            "Total: 0 subjects", size=13,
            color=ft.colors.GREY_500, weight=ft.FontWeight.W_500)

        # Add dialog to overlay
        if self.dialog not in self.page.overlay:
            self.page.overlay.append(self.dialog)
    
    def build(self):
        """Build the subjects view."""
        self.load_subjects()
        
        return ft.View(
            "/subjects",
            [
                # ── Navy top bar ───────────────────────────────────────
                ft.Container(
                    content=ft.Row([
                        ft.IconButton(
                            icon=ft.icons.ARROW_BACK_ROUNDED,
                            icon_color=ft.colors.WHITE, icon_size=20,
                            tooltip="Back to Dashboard",
                            on_click=lambda e: self.page.go("/dashboard"),
                        ),
                        ft.Container(width=4),
                        ft.Icon(ft.icons.MENU_BOOK_ROUNDED,
                                color=ft.colors.WHITE, size=20),
                        ft.Text("Subjects Management", size=17,
                                weight=ft.FontWeight.BOLD,
                                color=ft.colors.WHITE),
                        ft.Container(expand=True),
                        self._lbl_count_top,
                        ft.Container(width=8),
                        self._btn_refresh,
                    ], spacing=6,
                       vertical_alignment=ft.CrossAxisAlignment.CENTER),
                    bgcolor="#0F1E3A",
                    padding=ft.padding.symmetric(horizontal=16, vertical=10),
                    height=56,
                ),

                ft.Container(
                    content=ft.Column([
                        # ── Toolbar ────────────────────────────────────
                        ft.Container(
                            content=ft.Row([
                                ft.Container(
                                    content=self.search_field,
                                    expand=True,
                                    bgcolor=ft.colors.WHITE,
                                    border_radius=10,
                                    padding=ft.padding.symmetric(
                                        horizontal=4, vertical=2),
                                    border=ft.border.all(1, ft.colors.GREY_200),
                                ),
                                ft.Container(width=8),
                                ft.ElevatedButton(
                                    "+ Add Subject",
                                    visible=_can_edit(self._role),
                                    on_click=self.open_add_dialog,
                                    style=ft.ButtonStyle(
                                        bgcolor={"": "#2563EB", "hovered": "#1E40AF"},
                                        color=ft.colors.WHITE,
                                        shape=ft.RoundedRectangleBorder(radius=8),
                                        elevation={"": 0, "hovered": 2},
                                        padding=ft.padding.symmetric(
                                            horizontal=18, vertical=12),
                                    ),
                                ),
                                ft.Container(width=8),
                                ft.OutlinedButton(
                                    "Bulk Import",
                                    icon=ft.icons.UPLOAD_FILE_ROUNDED,
                                    visible=_can_edit(self._role),
                                    on_click=self.bulk_import,
                                    style=ft.ButtonStyle(
                                        shape=ft.RoundedRectangleBorder(radius=8),
                                        side=ft.BorderSide(1, "#2563EB"),
                                        color="#2563EB",
                                        padding=ft.padding.symmetric(
                                            horizontal=16, vertical=12),
                                    ),
                                ),
                            ], spacing=0),
                            bgcolor=ft.colors.WHITE,
                            border_radius=12,
                            padding=ft.padding.symmetric(horizontal=16, vertical=10),
                            shadow=ft.BoxShadow(
                                spread_radius=0, blur_radius=8,
                                color=ft.colors.with_opacity(0.06, ft.colors.BLACK),
                                offset=ft.Offset(0, 2)),
                        ),

                        ft.Container(height=12),

                        # ── Category tabs ──────────────────────────────
                        ft.Container(
                            content=self.category_tabs,
                            bgcolor=ft.colors.WHITE,
                            border_radius=12,
                            padding=ft.padding.symmetric(horizontal=12, vertical=4),
                            shadow=ft.BoxShadow(
                                spread_radius=0, blur_radius=8,
                                color=ft.colors.with_opacity(0.06, ft.colors.BLACK),
                                offset=ft.Offset(0, 2)),
                        ),

                        ft.Container(height=12),

                        # ── Bulk-select bar ────────────────────────────
                        self._sel_bar,

                        # ── Table card (no expand=True — fixes blank table) ──
                        ft.Container(
                            content=ft.Column([
                                ft.Row([
                                    self._lbl_count_table,
                                ], alignment=ft.MainAxisAlignment.END),
                                ft.Divider(height=1, color=ft.colors.GREY_100),
                                ft.Row(
                                    [self.subjects_table],
                                    alignment=ft.MainAxisAlignment.START,
                                    scroll=ft.ScrollMode.AUTO,
                                ),
                            ]),
                            bgcolor=ft.colors.WHITE,
                            border_radius=12,
                            padding=ft.padding.symmetric(horizontal=20, vertical=14),
                            shadow=ft.BoxShadow(
                                spread_radius=0, blur_radius=10,
                                color=ft.colors.with_opacity(0.07, ft.colors.BLACK),
                                offset=ft.Offset(0, 3)),
                        ),
                    ],
                    horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
                    scroll=ft.ScrollMode.AUTO,
                    spacing=0),
                    padding=20,
                    bgcolor="#F8FAFC",
                    expand=True,
                )
            ],
            padding=0,
            spacing=0,
        )
    
    def load_subjects(self):
        """Load subjects from database."""
        try:
            subjects = self.app.local_db.execute(
                "SELECT * FROM subjects ORDER BY code"
            )
            self.all_subjects = subjects if subjects else []
            self.selected_ids.clear()
            self.display_subjects()
            print(f"Loaded {len(self.all_subjects)} subjects")
        except Exception as e:
            print(f"Error loading subjects: {e}")
            traceback.print_exc()
            self.all_subjects = []
    
    def display_subjects(self, subjects=None):
        """Display subjects in table."""
        if subjects is None:
            subjects = self.all_subjects

        rows = []
        for subject in subjects:
            sid = subject.get('id', '')
            def _toggle(e, sid=sid):
                if sid in self.selected_ids:
                    self.selected_ids.discard(sid)
                else:
                    self.selected_ids.add(sid)
                n = len(self.selected_ids)
                self._sel_bar.visible = n > 0
                self._sel_count_lbl.value = f"{n} subject(s) selected"
                self.page.update()
            rows.append(
                ft.DataRow(
                    cells=[
                        ft.DataCell(ft.Checkbox(value=sid in self.selected_ids, on_change=_toggle, visible=_can_edit(self._role))),
                        ft.DataCell(ft.Text(subject.get('code', ''), size=14)),
                        ft.DataCell(ft.Text(subject.get('name', ''), size=14)),
                        ft.DataCell(
                            ft.Container(
                                content=ft.Text(
                                    subject.get('category', ''),
                                    size=12,
                                    color=ft.colors.WHITE
                                ),
                                padding=ft.padding.symmetric(horizontal=5, vertical=2),
                                border_radius=10,
                                bgcolor=self._get_category_color(subject.get('category', ''))
                            )
                        ),
                        ft.DataCell(ft.Text(str(subject.get('periods_per_week', 0)), size=14)),
                        ft.DataCell(ft.Text(f"{subject.get('duration_minutes', 40)} min", size=14)),
                        ft.DataCell(
                            ft.Icon(
                                ft.icons.CHECK_CIRCLE if subject.get('requires_lab', 0) else ft.icons.CANCEL,
                                color=ft.colors.GREEN if subject.get('requires_lab', 0) else ft.colors.RED
                            )
                        ),
                        ft.DataCell(
                            ft.Row([
                                ft.IconButton(
                                    icon=ft.icons.EDIT,
                                    icon_size=20,
                                    visible=_can_edit(self._role),
                                    on_click=lambda e, s=subject: self.open_edit_dialog(s)
                                ),
                                ft.IconButton(
                                    icon=ft.icons.DELETE,
                                    icon_size=20,
                                    visible=_can_edit(self._role),
                                    icon_color=ft.colors.RED_400,
                                    on_click=lambda e, s=subject: self.delete_subject(s)
                                ),
                                ft.Tooltip(
                                    message="View only",
                                    content=ft.Icon(ft.icons.VISIBILITY_OUTLINED,
                                                    size=18, color=ft.colors.GREY_400,
                                                    visible=not _can_edit(self._role))),
                            ])
                        ),
                    ]
                )
            )

        self.subjects_table.rows = rows
        n_total    = len(self.all_subjects)
        n_shown    = len(subjects)
        self._lbl_count_top.value = f"{n_total} subject{'s' if n_total != 1 else ''}"
        self._lbl_count_table.value = (
            f"Total: {n_total} subject{'s' if n_total != 1 else ''}"
            if n_shown == n_total else
            f"Showing: {n_shown} of {n_total} subject{'s' if n_total != 1 else ''}"
        )
        self.page.update()
    
    def _get_category_color(self, category):
        """Get color for category badge."""
        colors = {
            'Core':      "#2563EB",
            'Elective':  "#16A34A",
            'Practical': "#D97706",
        }
        return colors.get(category, "#64748B")
    
    def search_subjects(self, e):
        """Filter subjects by search term."""
        search = self.search_field.value.lower() if self.search_field.value else ""
        if not search:
            self.display_subjects()
            return
        
        filtered = [
            s for s in self.all_subjects
            if search in s.get('code', '').lower()
            or search in s.get('name', '').lower()
        ]
        self.display_subjects(filtered)
    
    def filter_by_category(self, e):
        """Filter subjects by category."""
        tab_index = self.category_tabs.selected_index
        categories = ["All", "Core", "Elective", "Practical"]
        selected = categories[tab_index]
        
        if selected == "All":
            self.display_subjects()
        else:
            filtered = [s for s in self.all_subjects if s.get('category') == selected]
            self.display_subjects(filtered)
    
    def open_add_dialog(self, e):
        """Open dialog to add new subject."""
        if not _can_edit(self._role): return
        self.dialog.title = ft.Text("Add Subject")
        self.current_subject_id = None
        
        content = self.dialog.content.content
        content.controls[0].value = ""
        content.controls[1].value = ""
        content.controls[2].value = "Core"
        content.controls[3].value = "3"
        content.controls[4].value = "40"
        content.controls[5].value = False
        
        self.dialog.open = True
        self.page.update()
    
    def open_edit_dialog(self, subject):
        """Open dialog to edit subject."""
        if not _can_edit(self._role): return
        self.dialog.title = ft.Text("Edit Subject")
        self.current_subject_id = subject['id']
        
        content = self.dialog.content.content
        content.controls[0].value = subject.get('code', '')
        content.controls[1].value = subject.get('name', '')
        content.controls[2].value = subject.get('category', 'Core')
        content.controls[3].value = str(subject.get('periods_per_week', 3))
        content.controls[4].value = str(subject.get('duration_minutes', 40))
        content.controls[5].value = bool(subject.get('requires_lab', False))
        
        self.dialog.open = True
        self.page.update()
    
    # ══════════════════════════════════════════════════════════════════════
    # BUTTON ANIMATION HELPERS
    # ══════════════════════════════════════════════════════════════════════
    async def _animate_elevated_btn(self, btn, save_fn, e):
        """Spinner → save logic → success flash → reset."""
        import asyncio
        orig_text  = btn.text
        orig_icon  = btn.icon
        # Saving state
        btn.text     = "Saving..."
        btn.icon     = ft.icons.HOURGLASS_TOP_ROUNDED
        btn.disabled = True
        btn.style.bgcolor = {"": "#64748B", "hovered": "#475569"}
        self.page.update()
        await asyncio.sleep(0.15)
        # Run logic
        save_fn(e)
        # Success state
        btn.text     = "Saved ✓"
        btn.icon     = ft.icons.CHECK_CIRCLE_ROUNDED
        btn.disabled = False
        btn.style.bgcolor = {"": "#16A34A", "hovered": "#15803D"}
        self.page.update()
        await asyncio.sleep(1.4)
        # Reset
        btn.text     = orig_text
        btn.icon     = orig_icon
        btn.style.bgcolor = {"": "#2563EB", "hovered": "#1E40AF"}
        self.page.update()

    async def _save_subject_animated(self, e):
        await self._animate_elevated_btn(self._btn_save_subject, self.save_subject, e)

    def save_subject(self, e):
        """Save subject to database."""
        if not _can_edit(self._role): return
        content = self.dialog.content.content
        
        subject_data = {
            'code': content.controls[0].value.upper(),
            'name': content.controls[1].value,
            'category': content.controls[2].value,
            'periods_per_week': int(content.controls[3].value),
            'duration_minutes': int(content.controls[4].value),
            'requires_lab': 1 if content.controls[5].value else 0,
        }
        
        # Validate
        if not subject_data['code'] or not subject_data['name']:
            self.page.show_snack_bar(
                ft.SnackBar(content=ft.Text("Code and Name are required!"), bgcolor=ft.colors.RED)
            )
            return
        
        try:
            if self.current_subject_id:
                # Update existing
                self.app.local_db.update(
                    "subjects", self.current_subject_id, subject_data
                )
            else:
                # Insert new
                subject_data['id'] = str(uuid.uuid4())
                self.app.local_db.insert("subjects", subject_data)
            
            self.close_dialog(None)
            self.load_subjects()
            
            self.page.show_snack_bar(
                ft.SnackBar(content=ft.Text("Subject saved successfully!"), bgcolor=ft.colors.GREEN)
            )
        except Exception as e:
            print(f"Error saving subject: {e}")
            traceback.print_exc()
            self.page.show_snack_bar(
                ft.SnackBar(content=ft.Text(f"Error: {str(e)}"), bgcolor=ft.colors.RED)
            )
    
    def _clear_selection(self, e=None):
        self.selected_ids.clear()
        self._sel_bar.visible = False
        self._sel_count_lbl.value = "0 selected"
        self.load_subjects()

    def _bulk_delete_confirm(self, e):
        if not _can_edit(self._role): return
        n = len(self.selected_ids)
        if n == 0:
            self.page.show_snack_bar(ft.SnackBar(
                content=ft.Text("No subjects selected. Tick the checkboxes first."),
                bgcolor=ft.colors.ORANGE))
            return
        ids_snap = set(self.selected_ids)

        def _do(e):
            dlg.open = False
            self.page.update()
            deleted = 0
            for sid in ids_snap:
                try:
                    # Cascade: remove assignments first, then the subject
                    self.app.local_db.execute(
                        "DELETE FROM class_subjects WHERE subject_id=?", (sid,))
                    self.app.local_db.execute(
                        "DELETE FROM teacher_subjects WHERE subject_id=?", (sid,))
                    self.app.local_db.delete("subjects", sid)
                    deleted += 1
                except Exception as ex:
                    print(f"Bulk delete {sid}: {ex}")
            self.selected_ids.clear()
            self._sel_bar.visible = False
            self._sel_count_lbl.value = "0 selected"
            self.load_subjects()
            self.page.show_snack_bar(ft.SnackBar(
                content=ft.Text(f"Deleted {deleted} subject(s) successfully."),
                bgcolor=ft.colors.GREEN))

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text("Delete Selected Subjects", color=ft.colors.RED_700),
            content=ft.Text(
                f"Delete {n} subject(s)? Their class and teacher assignments will also be removed. This cannot be undone."),
            actions=[
                ft.TextButton("Cancel",
                    on_click=lambda e: [setattr(dlg, 'open', False), self.page.update()]),
                ft.ElevatedButton("Delete All",
                    icon=ft.icons.DELETE_FOREVER,
                    on_click=_do,
                    style=ft.ButtonStyle(
                        bgcolor=ft.colors.RED_700,
                        color=ft.colors.WHITE,
                        shape=ft.RoundedRectangleBorder(radius=8),
                    )),
            ],
        )
        self.page.overlay.append(dlg)
        dlg.open = True
        self.page.update()

    def delete_subject(self, subject):
        """Delete subject with confirmation."""
        if not _can_edit(self._role): return
        def confirm_delete(e):
            try:
                # Remove any class assignments first (cascade), then delete.
                # This mirrors how delete_class works — we never block deletion
                # just because the subject is assigned somewhere.
                self.app.local_db.execute(
                    "DELETE FROM class_subjects WHERE subject_id = ?",
                    (subject['id'],)
                )
                # Also remove teacher-subject links
                self.app.local_db.execute(
                    "DELETE FROM teacher_subjects WHERE subject_id = ?",
                    (subject['id'],)
                )
                self.app.local_db.delete("subjects", subject['id'])
                dialog.open = False
                self.load_subjects()
                self.page.show_snack_bar(
                    ft.SnackBar(content=ft.Text("Subject deleted successfully!"),
                                bgcolor=ft.colors.GREEN)
                )
                self.page.update()
            except Exception as ex:
                print(f"Error deleting subject: {ex}")
                self.page.show_snack_bar(
                    ft.SnackBar(content=ft.Text(f"Error deleting subject: {ex}"),
                                bgcolor=ft.colors.RED)
                )
                self.page.update()
        
        dialog = ft.AlertDialog(
            title=ft.Text("Confirm Delete"),
            content=ft.Text(f"Are you sure you want to delete {subject.get('code')} - {subject.get('name')}?"),
            actions=[
                ft.TextButton("Cancel", on_click=lambda e: [setattr(dialog, 'open', False), self.page.update()]),
                ft.ElevatedButton("Delete", on_click=confirm_delete, color=ft.colors.RED),
            ]
        )
        
        if dialog not in self.page.overlay:
            self.page.overlay.append(dialog)
        dialog.open = True
        self.page.update()
    
    def bulk_import(self, e):
        """Open file picker — accepts CSV and XLSX."""
        if not _can_edit(self._role): return
        def on_file_picked(result: ft.FilePickerResultEvent):
            if result.files:
                self.process_bulk_import(result.files[0].path)

        picker = ft.FilePicker(on_result=on_file_picked)
        self.page.overlay.append(picker)
        self.page.update()
        picker.pick_files(
            dialog_title="Select Subjects file (CSV or Excel)",
            allow_multiple=False,
            allowed_extensions=["csv", "xlsx", "xls"],
        )

    def _read_import_rows(self, file_path):
        """Return list of dicts from CSV or XLSX. All keys lowercased & stripped."""
        import csv, os
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            raw_headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            # normalise: lowercase + strip, blank columns get positional name
            headers = [
                str(h).lower().strip() if h else f'col_{i}'
                for i, h in enumerate(raw_headers)
            ]
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if all(v is None for v in row):
                    continue  # skip fully empty rows
                rows.append({
                    headers[i]: (str(v).strip() if v is not None else '')
                    for i, v in enumerate(row)
                })
            return rows
        else:
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f)
                # normalise CSV headers too
                rows = []
                for raw_row in reader:
                    rows.append({k.lower().strip(): (v.strip() if v else '') for k, v in raw_row.items()})
                return rows

    def _col(self, row, *aliases):
        """Get first matching alias from a row dict (case already lowered)."""
        for a in aliases:
            v = row.get(a.lower().strip(), '')
            if v:
                return v
        return ''

    def process_bulk_import(self, file_path):
        """Process bulk import from CSV or XLSX.

        Accepted column names (case-insensitive):
          Subject Code  : code, subject_code, subject code
          Subject Name  : name, subject_name, subject name
          Category      : category
          Periods/Week  : periods/week, periods_per_week, periods per week
          Duration      : duration, duration_minutes
          Requires Lab  : requires_lab, requires lab, lab
        """
        try:
            rows = self._read_import_rows(file_path)

            if not rows:
                self.page.show_snack_bar(ft.SnackBar(
                    content=ft.Text("File is empty or could not be read."),
                    bgcolor=ft.colors.ORANGE))
                return

            detected = list(rows[0].keys())
            print(f"[Subjects Import] Detected headers: {detected}")

            def _int(val, default):
                try: return int(str(val).strip().split('.')[0] or str(default))
                except: return default

            count = 0
            skipped = 0
            skip_reasons = []

            for i, row in enumerate(rows, start=2):
                code = (row.get('code') or row.get('subject_code') or
                        row.get('subject code') or '').strip().upper()
                name = (row.get('name') or row.get('subject_name') or
                        row.get('subject name') or '').strip()

                if not code and not name:
                    skipped += 1; continue
                if not code:
                    skip_reasons.append(f"Row {i}: missing code"); skipped += 1; continue
                if not name:
                    skip_reasons.append(f"Row {i}: missing name"); skipped += 1; continue

                lab_raw = (row.get('requires_lab') or row.get('requires lab') or
                           row.get('lab') or 'false').lower().strip()
                subject_data = {
                    'id':               str(uuid.uuid4()),
                    'code':             code,
                    'name':             name,
                    'category':         (row.get('category') or 'Core').strip() or 'Core',
                    'periods_per_week': _int(row.get('periods/week') or row.get('periods_per_week') or
                                             row.get('periods per week'), 3),
                    'duration_minutes': _int(row.get('duration') or row.get('duration_minutes'), 40),
                    'requires_lab':     1 if lab_raw in ('true', '1', 'yes') else 0,
                }
                try:
                    self.app.local_db.insert("subjects", subject_data)
                    count += 1
                except Exception as row_ex:
                    skip_reasons.append(f"Row {i} ({code}): DB error — {row_ex}")
                    skipped += 1

            for r in skip_reasons:
                print(f"[Subjects Import] Skipped: {r}")

            self.load_subjects()
            msg = f"Imported {count} subject(s) successfully!"
            if skipped:
                msg += f" ({skipped} row(s) skipped)"
            color = ft.colors.GREEN if count > 0 else ft.colors.ORANGE
            self.page.show_snack_bar(
                ft.SnackBar(content=ft.Text(msg), bgcolor=color))

            if count == 0 and skipped > 0:
                hint = ft.AlertDialog(
                    title=ft.Text("Import Failed — Column Mismatch",
                                  color=ft.colors.ORANGE_700),
                    content=ft.Container(
                        width=480,
                        content=ft.Column([
                            ft.Text("No rows could be imported. Detected headers:", size=13),
                            ft.Container(
                                content=ft.Text(", ".join(detected) or "(none)",
                                                size=12, color=ft.colors.BLUE_700,
                                                selectable=True),
                                bgcolor=ft.colors.BLUE_50, border_radius=8, padding=10),
                            ft.Text("Expected headers:", size=13,
                                    weight=ft.FontWeight.BOLD, color=ft.colors.GREY_700),
                            ft.Text(
                                "- code (or: subject_code)\n- name (or: subject_name)\n- category\n- periods/week (or: periods_per_week)\n- duration (or: duration_minutes)\n- requires_lab (or: requires lab)",
                                size=12, color=ft.colors.GREY_800),
                        ], spacing=10, tight=True),
                    ),
                    actions=[ft.TextButton(
                        "OK", on_click=lambda e: setattr(hint, 'open', False)
                              or self.page.update())],
                )
                self.page.overlay.append(hint)
                hint.open = True
                self.page.update()

        except Exception as ex:
            print(f"Import error: {ex}")
            traceback.print_exc()
            self.page.show_snack_bar(
                ft.SnackBar(content=ft.Text(f"Import error: {str(ex)}"),
                            bgcolor=ft.colors.RED))

    def close_dialog(self, e):
        """Close the dialog."""
        self.dialog.open = False
        self.page.update()
    
    async def refresh(self, e):
        """Refresh subjects list with spinner animation."""
        import asyncio
        btn = self._btn_refresh
        btn.icon       = ft.icons.HOURGLASS_TOP_ROUNDED
        btn.icon_color = ft.colors.YELLOW_200
        btn.disabled   = True
        self.page.update()
        await asyncio.sleep(0.15)
        self.load_subjects()
        btn.icon       = ft.icons.CHECK_ROUNDED
        btn.icon_color = "#4ADE80"
        btn.disabled   = False
        self.page.update()
        await asyncio.sleep(1.2)
        btn.icon       = ft.icons.REFRESH_ROUNDED
        btn.icon_color = ft.colors.WHITE
        self.page.update()
        self.page.show_snack_bar(
            ft.SnackBar(content=ft.Text("Subjects refreshed!"), bgcolor=ft.colors.GREEN)
        )